#!/usr/bin/env python3
"""Lichtgewicht pptx->PNG renderer (Pillow) om decks visueel te kunnen inspecteren."""
import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPI = 100
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONTB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_cache = {}

def font(size_pt, bold):
    key = (int(size_pt*100), bold)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONTB if bold else FONT, max(6, int(size_pt*PPI/72)))
    return _cache[key]

def emu_px(v): return int((v or 0)/914400*PPI)

def rgb(c):
    try: return (c[0], c[1], c[2])
    except Exception: return None

def shape_fill(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:  # solid
            return (sh.fill.fore_color.rgb[0], sh.fill.fore_color.rgb[1], sh.fill.fore_color.rgb[2])
    except Exception: pass
    return None

def shape_line(sh):
    try:
        c = sh.line.color.rgb
        w = sh.line.width
        if c is None: return None, 0
        return (c[0], c[1], c[2]), max(1, emu_px(w) if w else 1)
    except Exception:
        return None, 0

def wrap(draw, text, fnt, maxw):
    out = []
    for raw in text.split("\n"):
        if raw == "":
            out.append(""); continue
        words = raw.split(" ")
        line = ""
        for w in words:
            t = (line + " " + w).strip()
            if draw.textlength(t, font=fnt) <= maxw or not line:
                line = t
            else:
                out.append(line); line = w
        out.append(line)
    return out

def render_slide(slide, idx):
    W = emu_px(prs.slide_width); H = emu_px(prs.slide_height)
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        if sh.left is None:
            pass
        x, y, w, h = emu_px(sh.left), emu_px(sh.top), emu_px(sh.width), emu_px(sh.height)
        # shape geometry
        fill = shape_fill(sh)
        lc, lw = shape_line(sh)
        is_diamond = False
        try: is_diamond = (sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 'DIAMOND' in str(sh.auto_shape_type))
        except Exception: pass
        if fill or lc:
            if is_diamond:
                d.polygon([(x+w//2,y),(x+w,y+h//2),(x+w//2,y+h),(x,y+h//2)],
                          fill=fill, outline=lc)
            else:
                d.rectangle([x, y, x+w, y+h], fill=fill, outline=lc, width=lw or 1)
        # text
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            pad = 5
            # collect lines with their font
            blocks = []
            for para in tf.paragraphs:
                runs = para.runs
                if runs:
                    txt = "".join(r.text for r in runs)
                    f = runs[0].font
                else:
                    txt = para.text; f = para.font
                size = f.size.pt if f.size else 12
                bold = bool(f.bold)
                col = rgb(f.color.rgb) if (f.color and f.color.type is not None) else (0,0,0)
                if col is None: col = (0,0,0)
                al = para.alignment
                fnt = font(size, bold)
                for ln in wrap(d, txt, fnt, max(10, w-2*pad)):
                    lh = int(size*PPI/72*1.25)
                    blocks.append((ln, fnt, col, al, lh))
            total = sum(b[4] for b in blocks)
            anchor = tf.vertical_anchor
            if anchor == MSO_ANCHOR.MIDDLE: cy = y + (h-total)//2
            elif anchor == MSO_ANCHOR.BOTTOM: cy = y + h - total
            else: cy = y + pad
            for ln, fnt, col, al, lh in blocks:
                tw = d.textlength(ln, font=fnt)
                if al == PP_ALIGN.CENTER: tx = x + (w-tw)//2
                elif al == PP_ALIGN.RIGHT: tx = x + w - tw - pad
                else: tx = x + pad
                d.text((tx, cy), ln, font=fnt, fill=col)
                cy += lh
    return img

path = sys.argv[1]
which = [int(a) for a in sys.argv[2:]] if len(sys.argv) > 2 else None
prs = Presentation(path)
import os
outdir = os.path.join(os.path.dirname(os.path.abspath(path)) or ".", "preview")
os.makedirs(outdir, exist_ok=True)
for i, sl in enumerate(prs.slides, 1):
    if which and i not in which: continue
    render_slide(sl, i).save(f"{outdir}/s{i:02d}.png")
print("rendered ->", outdir)
