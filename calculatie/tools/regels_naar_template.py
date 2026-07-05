#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
regels_naar_template.py — vul het Master calculatie-template (Monarch-layout)
regel voor regel met alle offerteregels uit offerteregels.json.

Gebruik: python3 calculatie/tools/regels_naar_template.py <slug> [uitvoer.xlsx]

Breidt secties met te weinig rijen automatisch uit (rijen invoegen vóór de
somrij) en schuift alle formules en kruisverwijzingen (Summary, Procore,
Check quote amounts, totaalbedrag) mee. Schrijft daarna elke offerteregel als
eigen regel in tabblad "4 | Calculation".
"""
import json, re, sys
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "template" / "Master_calculatie_LEEG.xlsx"
CALC = "4 | Calculation"
EENHEID = {"st": "stuks", "maal": "post"}  # naar de Categorie-keuzelijst


def item_rows_of(ws):
    """Rijen met het patroon M=K{r}*L{r} — de itemrijen van het rekenblad."""
    return {r for r in range(13, ws.max_row + 1)
            if ws.cell(row=r, column=13).value == f"=K{r}*L{r}"}


def section_map(ws):
    """code -> dict(lo, hi_fysiek, sumrow). hi_fysiek = laatste itemrij vóór de somrij."""
    items = item_rows_of(ws)
    last_code, out = None, {}
    for r in range(13, ws.max_row + 1):
        for col in (1, 2):
            v = ws.cell(row=r, column=col).value
            if v is not None and re.fullmatch(r"\d{3}", str(v).strip()):
                last_code = str(v).strip()
        g = ws.cell(row=r, column=7).value
        if isinstance(g, str):
            m = re.fullmatch(r"=SUM\(G(\d+):G(\d+)\)", g)
            if m and last_code and int(m.group(2)) - int(m.group(1)) < 400 \
                    and last_code not in out:
                lo = int(m.group(1))
                hi = lo
                while hi + 1 in items and hi + 1 < r:
                    hi += 1
                out[last_code] = {"lo": lo, "hi": hi, "sum": r}
    return out


def shift_ref(formula, shift):
    """Schuif rijnummers in celverwijzingen van een formule (zonder sheet-prefix-check)."""
    def rep(m):
        return f"{m.group(1)}{shift(int(m.group(2)))}"
    return re.sub(r"(\$?[A-Z]{1,3}\$?)(\d+)", rep, formula)


def shift_xrefs(formula, shift):
    """Schuif alleen rijen in verwijzingen naar '4 | Calculation' (incl. bereik-einde)."""
    pat = re.compile(r"('4 \| Calculation'!)(\$?[A-Z]{1,3}\$?)(\d+)(:(\$?[A-Z]{1,3}\$?)(\d+))?")
    def rep(m):
        s = f"{m.group(1)}{m.group(2)}{shift(int(m.group(3)))}"
        if m.group(4):
            s += f":{m.group(5)}{shift(int(m.group(6)))}"
        return s
    return pat.sub(rep, formula)


def copy_row_style(ws, src, dst):
    for col in range(1, 23):
        s, d = ws.cell(row=src, column=col), ws.cell(row=dst, column=col)
        d.font, d.border = copy(s.font), copy(s.border)
        d.fill, d.number_format = copy(s.fill), s.number_format
        d.alignment = copy(s.alignment)


def expand_section(wb, code, needed):
    """Voeg rijen toe aan een sectie zodat er `needed` itemrijen zijn."""
    ws = wb[CALC]
    sec = section_map(ws)[code]
    have = sec["hi"] - sec["lo"] + 1
    extra = needed - have
    if extra <= 0:
        return 0
    at = sec["hi"] + 1                      # invoegen direct na de laatste itemrij
    ws.insert_rows(at, extra)
    shift = lambda r: r + extra if r >= at else r

    # 1. alle formules in het rekenblad meeschuiven (semantiek blijft gelijk)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = shift_ref(cell.value, shift)
    # 2. nieuwe rijen opmaken als itemrij
    tpl = sec["lo"]
    for i in range(extra):
        r = at + i
        copy_row_style(ws, tpl, r)
        prev_b = ws.cell(row=r - 1, column=2).value
        ws.cell(row=r, column=2).value = (prev_b + 1) if isinstance(prev_b, int) else 1
        ws.cell(row=r, column=4).value = "post"
        ws.cell(row=r, column=5).value = f"=K{r}"
        ws.cell(row=r, column=6).value = f"=CEILING(L{r}*$M$7,$M$8)"
        ws.cell(row=r, column=7).value = f"=ROUND(E{r}*F{r},0)"
        ws.cell(row=r, column=9).value = ws.cell(row=tpl, column=9).value
        ws.cell(row=r, column=11).value = 0
        ws.cell(row=r, column=12).value = 0
        ws.cell(row=r, column=13).value = f"=K{r}*L{r}"
    # 3. somrij van deze sectie over het volledige (nieuwe) bereik laten lopen
    new_sum = shift(sec["sum"])
    ws.cell(row=new_sum, column=7).value = f"=SUM(G{sec['lo']}:G{at + extra - 1})"
    ws.cell(row=new_sum, column=13).value = f"=SUM(M{sec['lo']}:M{at + extra - 1})"
    # 4. kruisverwijzingen in de overige tabbladen meeschuiven
    for name in wb.sheetnames:
        if name == CALC:
            continue
        for row in wb[name].iter_rows(min_row=1, max_row=wb[name].max_row):
            for cell in row:
                if isinstance(cell.value, str) and "'4 | Calculation'!" in cell.value:
                    cell.value = shift_xrefs(cell.value, shift)
    # 5. keuzelijsten (datavalidatie) meeschuiven
    for dv in ws.data_validations.dataValidation:
        new = []
        for rng in str(dv.sqref).split():
            def rep(m):
                return f"{m.group(1)}{shift(int(m.group(2)))}"
            new.append(re.sub(r"([A-Z]{1,3})(\d+)", rep, rng))
        dv.sqref = " ".join(new)
    return extra


def build(slug, out=None):
    src = ROOT / "projecten" / slug / "offerteregels.json"
    data = json.loads(src.read_text(encoding="utf-8"))
    out = Path(out) if out else ROOT / "projecten" / slug / "calculatie.xlsx"

    # regels per sectiecode tellen en template-capaciteit bepalen
    lines = [(doc, regel) for doc in data["documenten"] for regel in doc["regels"]]
    per_code = {}
    for _, regel in lines:
        per_code[str(regel["code"])] = per_code.get(str(regel["code"]), 0) + 1

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[CALC]
    for code, needed in sorted(per_code.items(), key=lambda x: -section_map(ws)[x[0]]["lo"]):
        added = expand_section(wb, code, needed)
        if added:
            print(f"sectie {code}: {added} rijen toegevoegd")

    # regels wegschrijven
    secs = section_map(ws)
    cursor = {c: secs[c]["lo"] for c in per_code}
    for doc, regel in lines:
        code = str(regel["code"])
        r = cursor[code]
        if r > secs[code]["hi"]:
            raise SystemExit(f"FOUT: sectie {code} alsnog vol op rij {r}")
        cursor[code] = r + 1
        note = f"{doc['type'].capitalize()} {doc['referentie']} | {doc['bron']}"
        if regel.get("opmerking"):
            note += f" | {regel['opmerking']}"
        if doc["type"] == "offerte":
            note = "OFFERTE (nog geen opdracht) — " + note
        ws.cell(row=r, column=3).value = regel["omschrijving"]
        ws.cell(row=r, column=4).value = EENHEID.get(regel["eenheid"], regel["eenheid"])
        ws.cell(row=r, column=5).value = f"=K{r}"
        ws.cell(row=r, column=6).value = f"=CEILING(L{r}*$M$7,$M$8)"
        ws.cell(row=r, column=7).value = f"=ROUND(E{r}*F{r},0)"
        ws.cell(row=r, column=9).value = doc["leverancier"]
        ws.cell(row=r, column=11).value = regel["aantal"]
        ws.cell(row=r, column=12).value = regel["prijs"]
        ws.cell(row=r, column=13).value = f"=K{r}*L{r}"
        ws.cell(row=r, column=15).value = note

    # projectgegevens + leveranciers in controle-tab
    bp = wb["2 | Basic principles"]
    bp["C7"] = data["project"]
    bp["C8"] = "Inkoopoverzicht leveranciers — regel voor regel (offertes & opdrachten)"
    bp["C9"] = f"Peildatum {data['peildatum']}"
    bp["C23"] = "12e verdieping, WTC Beursplein 37 Rotterdam"
    bp["E23"] = 445
    suppliers = sorted({d["leverancier"] for d in data["documenten"]})
    cq = wb["Check quote amounts"]
    for i in range(25):
        cq.cell(row=15 + i, column=2).value = suppliers[i] if i < len(suppliers) else "-"

    wb.calculation.fullCalcOnLoad = True   # Excel herberekent alles bij openen
    wb.save(out)
    n = len(lines)
    tot = round(sum(round(rg["aantal"] * rg["prijs"], 2) for _, rg in lines), 2)
    print(f"OK: {out} — {n} regels, EUR {tot:,.2f} excl. btw")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("gebruik: regels_naar_template.py <slug> [uitvoer.xlsx]")
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
