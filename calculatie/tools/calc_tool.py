#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
calc_tool.py — beheer projectcalculaties op basis van het C&W Master_calculatie-template.

Gebruik (vanuit de repo-root):
  python3 calculatie/tools/calc_tool.py init "Projectnaam" [--slug naam] [--m2 445] [--locatie "..."]
  python3 calculatie/tools/calc_tool.py add <slug> --json <bestand.json | ->  [--force]
  python3 calculatie/tools/calc_tool.py status <slug>
  python3 calculatie/tools/calc_tool.py list

JSON-formaat voor `add` (één document = één offerte/opdracht/factuur):
{
  "leverancier": "SPIE",
  "referentie": "2500211",            // PO-, offerte- of factuurnummer (dedupe-sleutel)
  "type": "opdracht",                 // opdracht | offerte | factuur
  "bron": "01 SPIE/25002 ... .pdf",   // bestandsnaam/pad van de bron-PDF
  "datum": "2025-06-16",
  "regels": [
    {"code": "230", "omschrijving": "Elektrische installatie", "bedrag": 26603.03,
     "aantal": 1, "opmerking": "post 0018"}
  ]
}
Bedragen altijd netto excl. btw. "code" is de kostencode van het template
(110..143, 210..260, 311..370, 410..450, 510..620, 710..730).
"""
import argparse, datetime, json, re, shutil, sys
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]          # .../calculatie
TEMPLATE = ROOT / "template" / "Master_calculatie_LEEG.xlsx"
PROJECTS = ROOT / "projecten"
CALC_SHEET = "4 | Calculation"


# ---------- structuur van het rekenblad ----------

def section_map(ws):
    """code -> (eerste_itemrij, laatste_itemrij). Afgeleid uit A/B-kopjes en SUM-rijen."""
    last_code, out = None, {}
    for r in range(13, ws.max_row + 1):
        a, b = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        for v in (a, b):
            if v is not None and re.fullmatch(r"\d{3}", str(v).strip()):
                last_code = str(v).strip()
        g = ws.cell(row=r, column=7).value
        if isinstance(g, str):
            m = re.fullmatch(r"=SUM\(G(\d+):G(\d+)\)", g)
            if m and last_code and last_code not in out:
                lo, hi = int(m.group(1)), int(m.group(2))
                if hi - lo < 400:                     # sla de eindtotaal-rij over
                    out[last_code] = (lo, hi)
    return out


def write_line(ws, r, tpl_row, desc, party, qty, price, note):
    ws.cell(row=r, column=3).value = desc
    ws.cell(row=r, column=4).value = "post"
    ws.cell(row=r, column=5).value = f"=K{r}"
    ws.cell(row=r, column=6).value = f"=CEILING(L{r}*$M$7,$M$8)"
    ws.cell(row=r, column=7).value = f"=ROUND(E{r}*F{r},0)"
    ws.cell(row=r, column=9).value = party
    ws.cell(row=r, column=11).value = qty
    ws.cell(row=r, column=12).value = price
    ws.cell(row=r, column=13).value = f"=K{r}*L{r}"
    ws.cell(row=r, column=15).value = note
    if ws.cell(row=r, column=2).value in (None, ""):
        prev = ws.cell(row=r - 1, column=2).value
        ws.cell(row=r, column=2).value = (prev + 1) if isinstance(prev, int) else 1
    for col in (2, 3, 4, 5, 6, 7, 9, 11, 12, 13, 15):
        s, d = ws.cell(row=tpl_row, column=col), ws.cell(row=r, column=col)
        d.font, d.border = copy(s.font), copy(s.border)
        d.fill, d.number_format = copy(s.fill), s.number_format
        d.alignment = copy(s.alignment)


def free_row(ws, lo, hi, taken):
    """Eerst rijen zonder omschrijving, daarna ongebruikte placeholder-rijen (K=0)."""
    empties = [r for r in range(lo, hi + 1) if r not in taken
               and ws.cell(row=r, column=3).value in (None, "")]
    zeros = [r for r in range(lo, hi + 1) if r not in taken
             and ws.cell(row=r, column=11).value in (None, 0)]
    for r in empties + zeros:
        return r
    raise SystemExit(f"FOUT: geen vrije rij meer in sectie {lo}-{hi}")


def sync_check_sheet(wb, suppliers):
    cq = wb["Check quote amounts"]
    for i in range(25):
        cq.cell(row=15 + i, column=2).value = suppliers[i] if i < len(suppliers) else "-"


# ---------- projectbestanden ----------

def paths(slug):
    d = PROJECTS / slug
    return d, d / "calculatie.xlsx", d / "register.json"


def load_register(reg_path):
    if reg_path.exists():
        return json.loads(reg_path.read_text(encoding="utf-8"))
    return {"project": "", "offertes": []}


# ---------- commando's ----------

def cmd_init(args):
    slug = args.slug or re.sub(r"[^a-z0-9]+", "-", args.naam.lower()).strip("-")
    d, xlsx, reg = paths(slug)
    if xlsx.exists():
        raise SystemExit(f"FOUT: project '{slug}' bestaat al ({xlsx})")
    d.mkdir(parents=True, exist_ok=True)
    shutil.copy(TEMPLATE, xlsx)
    wb = openpyxl.load_workbook(xlsx)
    bp = wb["2 | Basic principles"]
    bp["C7"] = args.naam
    bp["C8"] = "Inkoopoverzicht leveranciers (offertes & opdrachten)"
    bp["C9"] = "Versie 001"
    bp["F9"] = datetime.date.today()
    if args.locatie:
        bp["C23"] = args.locatie
    if args.m2:
        bp["E23"] = args.m2
    wb.calculation.fullCalcOnLoad = True   # Excel herberekent alles bij openen
    wb.save(xlsx)
    reg.write_text(json.dumps({"project": args.naam, "slug": slug, "offertes": []},
                              indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"OK: project '{args.naam}' aangemaakt in {d}")


def cmd_add(args):
    d, xlsx, reg_path = paths(args.slug)
    if not xlsx.exists():
        raise SystemExit(f"FOUT: onbekend project '{args.slug}'. Draai eerst init of zie `list`.")
    raw = sys.stdin.read() if args.json == "-" else Path(args.json).read_text(encoding="utf-8")
    doc = json.loads(raw)
    for k in ("leverancier", "referentie", "type", "regels"):
        if not doc.get(k):
            raise SystemExit(f"FOUT: veld '{k}' ontbreekt in JSON")
    reg = load_register(reg_path)
    key = (doc["leverancier"].lower(), str(doc["referentie"]).lower())
    if not args.force and any((o["leverancier"].lower(), str(o["referentie"]).lower()) == key
                              for o in reg["offertes"]):
        raise SystemExit(f"OVERGESLAGEN: {doc['leverancier']} {doc['referentie']} staat al in het "
                         f"register (gebruik --force om toch toe te voegen)")

    wb = openpyxl.load_workbook(xlsx)
    ws = wb[CALC_SHEET]
    secs = section_map(ws)
    taken, total = set(), 0.0
    for regel in doc["regels"]:
        code = str(regel["code"])
        if code not in secs:
            raise SystemExit(f"FOUT: onbekende kostencode '{code}'. Geldig: {', '.join(sorted(secs))}")
        lo, hi = secs[code]
        r = free_row(ws, lo, hi, taken)
        taken.add(r)
        note = f"{doc['type'].capitalize()} {doc['referentie']}"
        if doc.get("bron"):
            note += f" | {doc['bron']}"
        if regel.get("opmerking"):
            note += f" | {regel['opmerking']}"
        if doc["type"].lower() == "offerte":
            note = "OFFERTE (nog geen opdracht) — " + note
        write_line(ws, r, lo, regel["omschrijving"], doc["leverancier"],
                   regel.get("aantal", 1), regel["bedrag"], note)
        total += regel.get("aantal", 1) * regel["bedrag"]

    reg["offertes"].append({
        "leverancier": doc["leverancier"], "referentie": str(doc["referentie"]),
        "type": doc["type"], "bron": doc.get("bron", ""), "datum": doc.get("datum", ""),
        "bedrag": round(total, 2), "regels": len(doc["regels"]),
        "toegevoegd": datetime.datetime.now().isoformat(timespec="seconds"),
    })
    suppliers = sorted({o["leverancier"] for o in reg["offertes"]})
    sync_check_sheet(wb, suppliers)
    wb.calculation.fullCalcOnLoad = True   # Excel herberekent alles bij openen
    wb.save(xlsx)
    reg_path.write_text(json.dumps(reg, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"OK: {doc['leverancier']} {doc['type']} {doc['referentie']}: "
          f"{len(doc['regels'])} regel(s), EUR {total:,.2f} excl. btw toegevoegd aan {xlsx.name}")


def cmd_status(args):
    d, xlsx, reg_path = paths(args.slug)
    if not xlsx.exists():
        raise SystemExit(f"FOUT: onbekend project '{args.slug}'")
    reg = load_register(reg_path)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[CALC_SHEET]
    tot = {}
    for r in range(13, ws.max_row + 1):
        m = ws.cell(row=r, column=13).value
        if isinstance(m, str) and m == f"=K{r}*L{r}":
            k = ws.cell(row=r, column=11).value or 0
            l = ws.cell(row=r, column=12).value
            if k and isinstance(l, (int, float)) and l:
                p = ws.cell(row=r, column=9).value
                tot[p] = tot.get(p, 0) + k * l
    print(f"Project: {reg.get('project', args.slug)}  ({len(reg['offertes'])} documenten verwerkt)")
    for p, v in sorted(tot.items(), key=lambda x: -x[1]):
        print(f"  {p:<22s} EUR {v:>12,.2f}")
    print(f"  {'TOTAAL netto excl. btw':<22s} EUR {sum(tot.values()):>12,.2f}")
    offertes_open = [o for o in reg["offertes"] if o["type"].lower() == "offerte"]
    if offertes_open:
        print("Nog geen opdracht (offertestatus):")
        for o in offertes_open:
            print(f"  - {o['leverancier']} {o['referentie']} (EUR {o['bedrag']:,.2f})")


def cmd_list(_args):
    if not PROJECTS.exists():
        print("(geen projecten)")
        return
    for d in sorted(PROJECTS.iterdir()):
        reg_path = d / "register.json"
        if reg_path.exists():
            reg = json.loads(reg_path.read_text(encoding="utf-8"))
            print(f"{d.name:<20s} {reg.get('project',''):<35s} {len(reg['offertes'])} documenten")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    pi = sub.add_parser("init"); pi.add_argument("naam"); pi.add_argument("--slug")
    pi.add_argument("--m2", type=float); pi.add_argument("--locatie"); pi.set_defaults(f=cmd_init)
    pa = sub.add_parser("add"); pa.add_argument("slug"); pa.add_argument("--json", required=True)
    pa.add_argument("--force", action="store_true"); pa.set_defaults(f=cmd_add)
    ps = sub.add_parser("status"); ps.add_argument("slug"); ps.set_defaults(f=cmd_status)
    pl = sub.add_parser("list"); pl.set_defaults(f=cmd_list)
    args = p.parse_args()
    args.f(args)


if __name__ == "__main__":
    main()
